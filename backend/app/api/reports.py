import io
import csv
from datetime import datetime, date
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.database import get_db
from app.models.employee import Employee
from app.models.department import Department
from app.models.attendance import Attendance
from app.models.payroll import Payroll
from app.models.risk import RiskAssessment
from app.models.leave import Leave
from app.models.task import Task
from app.models.user import User, UserRole
from app.schemas.reports import (
    ReportsSummaryResponse,
    EmployeeReportItem,
    AttritionReportItem,
    PayrollReportItem,
    AttendanceReportItem
)
from app.core.deps import get_current_user, require_roles
from app.core.audit import log_audit_event

router = APIRouter(prefix="/reports", tags=["Reports & Export System (Module 19)"])

@router.get("/summary", response_model=ReportsSummaryResponse)
def get_reports_summary(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles([UserRole.ADMIN, UserRole.HR, UserRole.MANAGER]))
):
    """Module 19: Overall Enterprise HR Metrics & KPI Summary for Reports Suite"""
    employees = db.query(Employee).all()
    total_emp = len(employees)
    active_emp = sum(1 for e in employees if e.status == "Active")
    avg_sal = round(sum(e.salary for e in employees) / total_emp, 2) if total_emp > 0 else 0.0

    today = date.today()
    payrolls = db.query(Payroll).filter(Payroll.month == today.month, Payroll.year == today.year).all()
    if not payrolls:
        payrolls = db.query(Payroll).all()
    total_payroll = sum(p.net_salary for p in payrolls)

    attendances = db.query(Attendance).all()
    total_att = len(attendances)
    present_att = sum(1 for a in attendances if a.status in ["Present", "Late"])
    att_rate = round((present_att / total_att) * 100, 1) if total_att > 0 else 92.5

    # Attrition stats
    risks = db.query(RiskAssessment).all()
    avg_risk = round(sum(r.risk_score for r in risks) / len(risks), 1) if risks else 24.5
    high_risk_count = sum(1 for r in risks if r.risk_score >= 50.0)

    pending_leaves = db.query(Leave).filter(Leave.status == "Pending").count()
    active_tasks = db.query(Task).filter(Task.status.in_(["TODO", "IN_PROGRESS", "REVIEW"])).count()

    return ReportsSummaryResponse(
        total_employees=total_emp,
        active_employees=active_emp,
        avg_salary=avg_sal,
        total_monthly_payroll=round(total_payroll, 2),
        attendance_rate=att_rate,
        avg_attrition_risk=avg_risk,
        high_risk_employee_count=high_risk_count,
        pending_leaves_count=pending_leaves,
        active_tasks_count=active_tasks
    )

@router.get("/employees", response_model=List[EmployeeReportItem])
def get_employee_report(
    department_id: Optional[int] = Query(None),
    status: Optional[str] = Query(None),
    employment_type: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles([UserRole.ADMIN, UserRole.HR, UserRole.MANAGER]))
):
    """Module 19: Comprehensive Employee Directory & Workforce Roster Report"""
    query = db.query(Employee).options(joinedload(Employee.department))
    if department_id:
        query = query.filter(Employee.department_id == department_id)
    if status:
        query = query.filter(Employee.status == status)
    if employment_type:
        query = query.filter(Employee.employment_type == employment_type)

    employees = query.order_by(Employee.id.asc()).all()

    # Pre-fetch latest risk assessments
    latest_risks = {}
    for r in db.query(RiskAssessment).order_by(RiskAssessment.id.asc()).all():
        latest_risks[r.employee_id] = r.risk_level

    items = []
    for emp in employees:
        items.append(EmployeeReportItem(
            id=emp.id,
            employee_code=emp.employee_code,
            name=f"{emp.first_name} {emp.last_name}",
            email=emp.email,
            department=emp.department.name if emp.department else "Unassigned",
            job_title=emp.job_title,
            salary=emp.salary,
            employment_type=emp.employment_type,
            status=emp.status,
            hire_date=emp.hire_date.strftime("%Y-%m-%d") if emp.hire_date else None,
            risk_level=latest_risks.get(emp.id, "LOW")
        ))
    return items

@router.get("/attrition", response_model=List[AttritionReportItem])
def get_attrition_report(
    risk_level: Optional[str] = Query(None),
    department_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles([UserRole.ADMIN, UserRole.HR, UserRole.MANAGER]))
):
    """Module 19: Enterprise AI Attrition Risk Assessment Report"""
    query = db.query(Employee).options(joinedload(Employee.department))
    if department_id:
        query = query.filter(Employee.department_id == department_id)
    employees = query.all()

    items = []
    for emp in employees:
        latest = db.query(RiskAssessment).filter(
            RiskAssessment.employee_id == emp.id
        ).order_by(RiskAssessment.id.desc()).first()

        score = latest.risk_score if latest else (35.0 if emp.salary < 80000 else 18.0)
        level = latest.risk_level if latest else ("HIGH" if score >= 50 else ("MEDIUM" if score >= 30 else "LOW"))
        driver = latest.primary_driver if latest else "Standard Baseline"
        assessed_date = latest.assessed_at.strftime("%Y-%m-%d %H:%M") if latest else datetime.utcnow().strftime("%Y-%m-%d %H:%M")

        if risk_level and level != risk_level:
            continue

        tenure = 2.5
        if emp.hire_date:
            days = (datetime.utcnow() - emp.hire_date).days
            tenure = round(days / 365.25, 1)

        items.append(AttritionReportItem(
            employee_id=emp.id,
            employee_code=emp.employee_code,
            name=f"{emp.first_name} {emp.last_name}",
            department=emp.department.name if emp.department else "Unassigned",
            job_title=emp.job_title,
            risk_score=score,
            risk_level=level,
            primary_driver=driver,
            tenure_years=tenure,
            last_assessment_date=assessed_date
        ))

    return sorted(items, key=lambda x: x.risk_score, reverse=True)

@router.get("/payroll", response_model=List[PayrollReportItem])
def get_payroll_report(
    month: Optional[int] = Query(None),
    year: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles([UserRole.ADMIN, UserRole.HR]))
):
    """Module 19: Financial & Compensation Payroll Breakdown Report"""
    query = db.query(Payroll).options(
        joinedload(Payroll.employee).joinedload(Employee.department)
    )
    if month:
        query = query.filter(Payroll.month == month)
    if year:
        query = query.filter(Payroll.year == year)

    records = query.order_by(Payroll.id.desc()).all()
    items = []
    for p in records:
        emp_name = f"{p.employee.first_name} {p.employee.last_name}" if p.employee else "Employee"
        dept_name = p.employee.department.name if (p.employee and p.employee.department) else "Operations"
        items.append(PayrollReportItem(
            id=p.id,
            employee_name=emp_name,
            department=dept_name,
            month_year=f"{p.month:02d}/{p.year}",
            basic_salary=p.basic_salary,
            allowances=p.allowances,
            bonus=p.bonus,
            tax_deduction=p.tax_deduction,
            provident_fund=p.provident_fund,
            net_salary=p.net_salary,
            status=p.status
        ))
    return items

@router.get("/attendance", response_model=List[AttendanceReportItem])
def get_attendance_report(
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles([UserRole.ADMIN, UserRole.HR, UserRole.MANAGER]))
):
    """Module 19: Workforce Attendance, Working Hours & Overtime Compliance Report"""
    query = db.query(Attendance).options(
        joinedload(Attendance.employee).joinedload(Employee.department)
    )
    if status:
        query = query.filter(Attendance.status == status)
    if date_from:
        query = query.filter(Attendance.date >= datetime.strptime(date_from, "%Y-%m-%d").date())
    if date_to:
        query = query.filter(Attendance.date <= datetime.strptime(date_to, "%Y-%m-%d").date())

    records = query.order_by(Attendance.date.desc()).limit(100).all()
    items = []
    for a in records:
        emp_name = f"{a.employee.first_name} {a.employee.last_name}" if a.employee else "Employee"
        dept_name = a.employee.department.name if (a.employee and a.employee.department) else "Operations"
        items.append(AttendanceReportItem(
            employee_name=emp_name,
            department=dept_name,
            date=a.date.strftime("%Y-%m-%d"),
            check_in=a.check_in.strftime("%H:%M:%S") if a.check_in else None,
            check_out=a.check_out.strftime("%H:%M:%S") if a.check_out else None,
            shift_type=a.shift_type or "Regular",
            working_hours=a.working_hours or 0.0,
            overtime_hours=a.overtime_hours or 0.0,
            status=a.status
        ))
    return items

# --- EXPORT ENDPOINTS: CSV, EXCEL, PDF (Module 19) ---

@router.get("/export/csv")
def export_csv(
    report_type: str = Query("employees", pattern="^(employees|attrition|payroll|attendance)$"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles([UserRole.ADMIN, UserRole.HR, UserRole.MANAGER]))
):
    """Module 19: Export Report Data to Clean CSV Format"""
    output = io.StringIO()
    writer = csv.writer(output)

    if report_type == "employees":
        writer.writerow(["Employee Code", "Full Name", "Email", "Department", "Job Title", "Salary (Annual)", "Type", "Status"])
        for emp in db.query(Employee).options(joinedload(Employee.department)).all():
            writer.writerow([
                emp.employee_code,
                f"{emp.first_name} {emp.last_name}",
                emp.email,
                emp.department.name if emp.department else "Unassigned",
                emp.job_title,
                emp.salary,
                emp.employment_type,
                emp.status
            ])
    elif report_type == "attrition":
        writer.writerow(["Employee Code", "Name", "Department", "Job Title", "Risk Score (%)", "Risk Level", "Primary Driver"])
        for emp in db.query(Employee).options(joinedload(Employee.department)).all():
            r = db.query(RiskAssessment).filter(RiskAssessment.employee_id == emp.id).order_by(RiskAssessment.id.desc()).first()
            writer.writerow([
                emp.employee_code,
                f"{emp.first_name} {emp.last_name}",
                emp.department.name if emp.department else "Unassigned",
                emp.job_title,
                r.risk_score if r else 25.0,
                r.risk_level if r else "LOW",
                r.primary_driver if r else "Standard Workload"
            ])
    elif report_type == "payroll":
        writer.writerow(["Employee Name", "Department", "Month/Year", "Basic ($)", "Allowances ($)", "Bonus ($)", "Gross ($)", "Tax ($)", "PF ($)", "Net Salary ($)", "Status"])
        for p in db.query(Payroll).options(joinedload(Payroll.employee).joinedload(Employee.department)).all():
            emp_name = f"{p.employee.first_name} {p.employee.last_name}" if p.employee else "Employee"
            dept = p.employee.department.name if (p.employee and p.employee.department) else "N/A"
            writer.writerow([
                emp_name, dept, f"{p.month:02d}/{p.year}", p.basic_salary, p.allowances, p.bonus, p.gross_salary, p.tax_deduction, p.provident_fund, p.net_salary, p.status
            ])
    elif report_type == "attendance":
        writer.writerow(["Date", "Employee Name", "Department", "Shift", "Check-In", "Check-Out", "Hours", "Overtime", "Status"])
        for a in db.query(Attendance).options(joinedload(Attendance.employee).joinedload(Employee.department)).all():
            emp_name = f"{a.employee.first_name} {a.employee.last_name}" if a.employee else "Employee"
            dept = a.employee.department.name if (a.employee and a.employee.department) else "N/A"
            writer.writerow([
                a.date.strftime("%Y-%m-%d"),
                emp_name,
                dept,
                a.shift_type,
                a.check_in.strftime("%H:%M:%S") if a.check_in else "-",
                a.check_out.strftime("%H:%M:%S") if a.check_out else "-",
                a.working_hours,
                a.overtime_hours,
                a.status
            ])

    log_audit_event(
        db=db,
        user_id=current_user.id,
        user_name=current_user.full_name,
        action="EXPORT_CSV",
        module="REPORTS_SYSTEM",
        details=f"Exported {report_type.upper()} report to CSV."
    )

    output.seek(0)
    filename = f"workforce_ai_{report_type}_report_{date.today().strftime('%Y%m%d')}.csv"
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/export/excel")
def export_excel(
    report_type: str = Query("employees", pattern="^(employees|attrition|payroll|attendance)$"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles([UserRole.ADMIN, UserRole.HR, UserRole.MANAGER]))
):
    """Module 19: Export Formatted Enterprise Excel (.xlsx) Spreadsheet with Corporate Styling"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{report_type.capitalize()} Report"

    # Header styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    border_side = Side(style='thin', color="CBD5E1")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    if report_type == "employees":
        headers = ["Employee Code", "Full Name", "Email", "Department", "Job Title", "Salary ($)", "Type", "Status"]
        rows = [
            [
                e.employee_code,
                f"{e.first_name} {e.last_name}",
                e.email,
                e.department.name if e.department else "Unassigned",
                e.job_title,
                e.salary,
                e.employment_type,
                e.status
            ]
            for e in db.query(Employee).options(joinedload(Employee.department)).all()
        ]
    elif report_type == "attrition":
        headers = ["Employee Code", "Name", "Department", "Job Title", "Risk Score (%)", "Risk Level", "Primary Driver"]
        rows = []
        for e in db.query(Employee).options(joinedload(Employee.department)).all():
            r = db.query(RiskAssessment).filter(RiskAssessment.employee_id == e.id).order_by(RiskAssessment.id.desc()).first()
            rows.append([
                e.employee_code,
                f"{e.first_name} {e.last_name}",
                e.department.name if e.department else "Unassigned",
                e.job_title,
                r.risk_score if r else 25.0,
                r.risk_level if r else "LOW",
                r.primary_driver if r else "Standard Baseline"
            ])
    elif report_type == "payroll":
        headers = ["Employee Name", "Department", "Month/Year", "Basic ($)", "Gross ($)", "Tax ($)", "PF ($)", "Net Salary ($)", "Status"]
        rows = []
        for p in db.query(Payroll).options(joinedload(Payroll.employee).joinedload(Employee.department)).all():
            emp_name = f"{p.employee.first_name} {p.employee.last_name}" if p.employee else "Employee"
            dept = p.employee.department.name if (p.employee and p.employee.department) else "N/A"
            rows.append([
                emp_name, dept, f"{p.month:02d}/{p.year}", p.basic_salary, p.gross_salary, p.tax_deduction, p.provident_fund, p.net_salary, p.status
            ])
    elif report_type == "attendance":
        headers = ["Date", "Employee Name", "Department", "Shift", "Check-In", "Check-Out", "Hours", "Overtime", "Status"]
        rows = []
        for a in db.query(Attendance).options(joinedload(Attendance.employee).joinedload(Employee.department)).all():
            emp_name = f"{a.employee.first_name} {a.employee.last_name}" if a.employee else "Employee"
            dept = a.employee.department.name if (a.employee and a.employee.department) else "N/A"
            rows.append([
                a.date.strftime("%Y-%m-%d"),
                emp_name,
                dept,
                a.shift_type,
                a.check_in.strftime("%H:%M:%S") if a.check_in else "-",
                a.check_out.strftime("%H:%M:%S") if a.check_out else "-",
                a.working_hours,
                a.overtime_hours,
                a.status
            ])

    # Write headers
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border

    # Write data rows
    for r_idx, row_data in enumerate(rows, start=2):
        ws.append(row_data)
        for c_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = data_font
            cell.border = cell_border
            if isinstance(row_data[c_idx - 1], (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-adjust column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    log_audit_event(
        db=db,
        user_id=current_user.id,
        user_name=current_user.full_name,
        action="EXPORT_EXCEL",
        module="REPORTS_SYSTEM",
        details=f"Exported {report_type.upper()} report to Excel (.xlsx)."
    )

    filename = f"workforce_ai_{report_type}_report_{date.today().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/export/pdf")
def export_pdf_report(
    report_type: str = Query("employees", pattern="^(employees|attrition|payroll|attendance)$"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles([UserRole.ADMIN, UserRole.HR, UserRole.MANAGER]))
):
    """Module 19: Export Printable Corporate PDF Report using ReportLab"""
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(letter),
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=6
    )
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=15
    )

    elements = []
    title_text = f"WorkForce AI Pro — {report_type.capitalize()} Report"
    date_str = datetime.utcnow().strftime("%B %d, %Y - %H:%M UTC")
    elements.append(Paragraph(title_text, title_style))
    elements.append(Paragraph(f"Generated on {date_str} | Generated by: {current_user.full_name} ({current_user.role})", subtitle_style))
    elements.append(Spacer(1, 10))

    if report_type == "employees":
        table_data = [["Code", "Name", "Department", "Job Title", "Salary ($)", "Type", "Status"]]
        for e in db.query(Employee).options(joinedload(Employee.department)).all():
            table_data.append([
                e.employee_code,
                f"{e.first_name} {e.last_name}",
                e.department.name if e.department else "N/A",
                e.job_title[:24],
                f"${e.salary:,.0f}",
                e.employment_type,
                e.status
            ])
    elif report_type == "attrition":
        table_data = [["Code", "Name", "Department", "Job Title", "Risk Score", "Level", "Primary Driver"]]
        for e in db.query(Employee).options(joinedload(Employee.department)).all():
            r = db.query(RiskAssessment).filter(RiskAssessment.employee_id == e.id).order_by(RiskAssessment.id.desc()).first()
            table_data.append([
                e.employee_code,
                f"{e.first_name} {e.last_name}",
                e.department.name if e.department else "N/A",
                e.job_title[:22],
                f"{r.risk_score:.1f}%" if r else "25.0%",
                r.risk_level if r else "LOW",
                r.primary_driver[:26] if r else "Baseline"
            ])
    elif report_type == "payroll":
        table_data = [["Employee", "Department", "Month/Year", "Basic", "Gross", "Tax", "PF", "Net Salary", "Status"]]
        for p in db.query(Payroll).options(joinedload(Payroll.employee).joinedload(Employee.department)).all():
            emp_name = f"{p.employee.first_name} {p.employee.last_name}" if p.employee else "Employee"
            dept = p.employee.department.name if (p.employee and p.employee.department) else "N/A"
            table_data.append([
                emp_name, dept, f"{p.month:02d}/{p.year}", f"${p.basic_salary:,.0f}", f"${p.gross_salary:,.0f}",
                f"${p.tax_deduction:,.0f}", f"${p.provident_fund:,.0f}", f"${p.net_salary:,.0f}", p.status
            ])
    elif report_type == "attendance":
        table_data = [["Date", "Employee", "Department", "Shift", "Check-In", "Check-Out", "Hours", "OT", "Status"]]
        for a in db.query(Attendance).options(joinedload(Attendance.employee).joinedload(Employee.department)).limit(40).all():
            emp_name = f"{a.employee.first_name} {a.employee.last_name}" if a.employee else "Employee"
            dept = a.employee.department.name if (a.employee and a.employee.department) else "N/A"
            table_data.append([
                a.date.strftime("%Y-%m-%d"),
                emp_name,
                dept,
                a.shift_type[:15] if a.shift_type else "Regular",
                a.check_in.strftime("%H:%M") if a.check_in else "-",
                a.check_out.strftime("%H:%M") if a.check_out else "-",
                f"{a.working_hours:.1f}",
                f"{a.overtime_hours:.1f}",
                a.status
            ])

    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('TOPPADDING', (0, 0), (-1, 0), 6),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F8FAFC')),
        ('TEXTCOLOR', (0, 1), (-1, -1), colors.HexColor('#1E293B')),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
    ]))

    elements.append(t)
    doc.build(elements)
    buf.seek(0)

    log_audit_event(
        db=db,
        user_id=current_user.id,
        user_name=current_user.full_name,
        action="EXPORT_PDF",
        module="REPORTS_SYSTEM",
        details=f"Exported {report_type.upper()} report to PDF."
    )

    filename = f"workforce_ai_{report_type}_report_{date.today().strftime('%Y%m%d')}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
